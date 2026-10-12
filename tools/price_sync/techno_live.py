# -*- coding: utf-8 -*-
"""Прайс Техносистемса без ручной правки `technosystems_price.csv`.

Владелец 12.10.2026: «хочу через токен — чтобы при свежем прайсе сайт обновлялся сам».
Источники, от свежего к общему:

  * посты их Telegram-канала «Техносистемс» — их читает личный Telegram владельца
    (сессия на сервере, контейнер telegram-personal-mcp; см. techno_tg_fetch.py).
    Текстовый прайс разбирается; новый пост перебивает таблицу, пока та не изменится.
    Прайс картинкой не разбираем — о нём приходит строчка в итоговом сообщении;
  * их Google-таблица «ПРАЙСЫ» с сайта техносистемс.рф: вкладка цен и вкладка
    складских остатков в штуках (с датой среза).

`technosystems_price.csv` остаётся реестром позиций: бренд, серия, тип — то, по чему
строятся карточки. Отсюда в нём обновляются только цены и наличие; позиции, которых в
реестре нет, дописываются (карточка появится, когда будут описание и фото).

    python -m tools.price_sync.techno_live            # обновить CSV, при новостях — в Telegram
    python -m tools.price_sync.techno_live --dry-run  # только показать
    python -m tools.price_sync.techno_live --last-id  # номер последнего прочитанного поста
"""
from __future__ import annotations

import argparse
import base64
import csv
import hashlib
import io
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

import requests

from tools.price_sync import sources_techno
from tools.price_sync.sync import DEFAULT_OUT

SHEET_KEY = "2PACX-1vSqQYrwpQuH3BNen7dlp0ays4nYBAMnwHnyMH1JZ4JEJ_IAAvUUGXvLCf5zmWjvgWEekT-hqwn0FP66"
SHEET_URL = "https://docs.google.com/spreadsheets/d/e/%s/pub?gid=%s&single=true&output=csv"
PRICE_GID, STOCK_GID = "0", "384325198"
#: Канал «Техносистемс» в Telegram (id канала без -100).
CHANNEL_ID = 2177370047

STATE_FILE = DEFAULT_OUT / "techno_state.json"
POSTS_FILE = DEFAULT_OUT / "techno_tg.json"
COLUMNS = ("бренд", "серия", "модель", "категория", "цена_1", "цена_2", "наличие", "примечание")

BRANDS = ("EUROHOFF", "DENKO", "DANTEX", "DAICHI", "MARSA", "DAHATSU", "YUETU", "YNOVIK")
_CYR = re.compile(r"[А-Яа-яЁё]")
_INV = re.compile(r"(?i)INV(?:ERTER|ERTOR|ERT)?\b")
_PRICE = re.compile(r"\d{1,3}(?:[  ]\d{3})+(?!\d)|\d+")
_LATIN_DIGIT = re.compile(r"(?=.*[A-Za-z])(?=.*\d)")


# ── Разбор ────────────────────────────────────────────────────────────────────


def prices_of(text: str) -> list[float]:
    """«18650 /  17850», «43 300 39 200», «72000 (GREE)», «25 900 ₽» → числа."""
    values = []
    for token in _PRICE.findall(text or ""):
        value = float(re.sub(r"\s", "", token))
        if value >= 100:
            values.append(value)
    return values


def model_key(text: str) -> tuple[str, bool]:
    """Код модели для сопоставления и признак «инвертор».

    «Denko DU-07 inverter» → (DU07, True); «RK-07SCDG/07SCDGE» → RK07SCDG (второй блок
    после «/» отбрасываем); «MULT-18/2» → MULT182 (цифры после «/» — часть кода);
    русские слова («нап пот», «нар блок», «шт.») и IN/OUT — не код.
    """
    upper = (text or "").upper()
    inverter = bool(_INV.search(upper))
    upper = _INV.sub(" ", upper)
    parts = upper.split("/")
    head = parts[0]
    tail = parts[1].split()[0] if len(parts) > 1 and parts[1].split() else ""
    if re.fullmatch(r"\d{1,2}|[A-Z]{1,2}\d", tail):
        head += tail          # «MULT-18/2»; «EUD100W1/T1» (кассета) ≠ «EUD100W1/ZD1»
    head = re.sub(r"(?<=\d)IN\b", "", head)   # «YEL-07IN/OUT» = «YEL-07 IN/YEL-07 OUT»
    tokens = [t for t in re.split(r"\s+", head) if t and not _CYR.search(t)]
    tokens = [t for t in tokens if t not in BRANDS and t not in ("IN", "OUT")]
    key = re.sub(r"[^A-Z0-9]", "", "".join(tokens))
    return key, inverter


def _is_model(text: str) -> bool:
    return bool(_LATIN_DIGIT.match(text or "")) and ".." not in (text or "")


def parse_sheet(text: str) -> list[dict]:
    """Вкладка цен: «Марка, Модель, Цена, НАЛИЧИЕ». Марка — заголовок секции, строки ниже
    её наследуют; строки без кода в «Модели» — подзаголовки серий."""
    entries, section, sub = [], "", ""
    reader = csv.reader(io.StringIO(text))
    next(reader, None)
    for cells in reader:
        cells = (cells + ["", "", "", ""])[:4]
        brand_cell, model, price, avail = (c.strip() for c in cells)
        model = re.sub(r"\s+", " ", model)
        if brand_cell:
            section, sub = re.sub(r"\s+", " ", brand_cell), ""
        if not model or model.lower() == "модель":
            continue
        if not _is_model(model) or not model_key(model)[0]:
            sub = model
            continue
        entries.append({"section": section, "sub": sub, "model": model,
                        "hint": bool(_INV.search(section + " " + sub)),
                        "prices": prices_of(price), "in_stock": avail.lower().startswith("в налич")})
    return entries


def parse_stock(text: str) -> tuple[str, list[dict]]:
    """Вкладка остатков: дата среза и строки «Модель (шт.), количество»."""
    date, items = "", []
    for cells in csv.reader(io.StringIO(text)):
        first = (cells[0] if cells else "").strip()
        found = re.fullmatch(r"(\d{2})\.(\d{2})\.(\d{2})", first)
        if found and not date:
            day, month, year = found.groups()
            date = "20%s-%s-%s" % (year, month, day)
            continue
        if "(шт.)" not in first:
            continue
        qty = cells[1].strip() if len(cells) > 1 else ""
        items.append({"model": first.replace("(шт.)", "").strip(),
                      "qty": int(float(qty.replace(",", "."))) if re.fullmatch(r"[\d.,]+", qty) else 0})
    return date, items


_SEPARATOR = re.compile(r"\s+[—–]\s*|\s*[—–]\s+|\s+-\s+|\s{2,}")


def parse_post(text: str) -> list[dict]:
    """Текстовый прайс из канала: «ICE20 — 25 900 ₽ ✅ | ICE25 — 28 450 ₽ ✅».

    ✅ — на складе. «уточняйте» / «под заказ» — позиция без цены. Диапазоны «DNP-36A..60A»
    пропускаем: какой модели какая цена, из них не понять.
    """
    entries, header = [], ""
    for line in (text or "").splitlines():
        for segment in line.split("|"):
            parts = _SEPARATOR.split(segment.strip(), maxsplit=1)
            if len(parts) != 2:
                if segment.strip():
                    header = segment.strip()      # «❄️ DAICHI ICE (-15°C Inv)» — заголовок
                continue
            model, rest = parts
            model = model.rsplit(":", 1)[-1].strip()
            model = re.sub(r"\s+[А-ЯЁ][А-ЯЁ\-]+$", "", model)   # «… КАНАЛЬНЫЙ», «… КАССЕТА»
            if not _is_model(model) or not model_key(model)[0]:
                continue
            entries.append({"model": model, "prices": prices_of(rest), "in_stock": "✅" in rest,
                            "hint": bool(_INV.search(header))})
    return entries


def parse_table_rows(rows: list[list]) -> list[dict]:
    """Прайс файлом (xlsx): строка с кодом модели и числами — позиция."""
    entries = []
    for row in rows:
        cells = [str(c).strip() for c in row if c not in (None, "")]
        model = next((c for c in cells if _is_model(c) and model_key(c)[0]), None)
        if not model:
            continue
        numbers = [p for c in cells if c != model for p in prices_of(c)]
        entries.append({"model": model, "prices": numbers[:2], "in_stock": False})
    return entries


# ── Сопоставление с реестром ─────────────────────────────────────────────────


def _find(key: str, inverter: bool, items: list[dict], multi_ok: bool = True,
          strict: bool = True, min_tail: int = 3) -> list[dict]:
    """Позиции источника для кода реестра: точное совпадение, иначе код источника
    длиннее нашего («ICE20» ↔ «ICE20AVQS1R2»). Хвост в 1–2 знака — уже другой товар
    (RK-07SCDG ↔ RK-07SCDGI). Метка «inv» у источника без нашей — не помеха
    («EM -07i inv» = «EM-07I»), а наш инвертор чужой on/off не берёт."""
    if not key:
        return []

    def usable(item):
        if not multi_ok and "мульти" in item["model"].lower():
            return False
        # Нашему инвертору чужой on/off не годится — когда у модели есть on/off-двойник
        # с тем же кодом (DU-07 / DU-07 inv). Без двойника метки «inv» в источнике может
        # не быть вовсе («EUD100W1/T1 касс»).
        return not (strict and inverter and not item["_inv"])

    exact = [i for i in items if i["_key"] == key and usable(i)]
    if exact:
        same = [i for i in exact if i["_inv"] == inverter]
        return same or exact
    if len(key) < 4:
        return []
    return [i for i in items if i["_key"].startswith(key) and len(i["_key"]) - len(key) >= min_tail
            and not i["_key"][len(key)] == "I" and usable(i)]


def _keyed(items: list[dict]) -> list[dict]:
    for item in items:
        item["_key"], item["_inv"] = model_key(item["model"])
        item["_inv"] = item["_inv"] or bool(item.get("hint"))
    return items


def _money(value: float) -> str:
    return str(int(value)) if value == int(value) else str(value)


def _category(text: str) -> str:
    low = text.lower()
    for marks, name in ((("касс",), "кассета"), (("нап",), "напольно-потолочный"),
                        (("кан",), "канальный"), (("колон",), "колонный"),
                        (("нар блок", "нар. блок"), "мультисплит наружный"),
                        (("внут блок", "вн.блок"), "мультисплит внутренний")):
        if any(mark in low for mark in marks):
            return name
    return "настенный"


def _brand(entry: dict) -> str:
    for text in (entry["model"], entry.get("section", ""), entry.get("sub", "")):
        for brand in BRANDS:
            if brand in text.upper():
                return brand
    return ""


def _new_row(entry: dict) -> dict:
    brand = _brand(entry)
    model = entry["model"]
    if model.upper().startswith(brand + " "):
        model = model[len(brand) + 1:]
    series = entry.get("sub") or entry.get("section") or ""
    if model_key(model)[1] and "inv" not in series.lower():
        series = (series + " inverter").strip()
    prices = entry["prices"]
    return {"бренд": brand, "серия": series, "модель": model,
            "категория": _category(entry["model"] + " " + series),
            "цена_1": _money(prices[0]) if prices else "",
            "цена_2": _money(prices[1]) if len(prices) > 1 else "",
            "наличие": "есть" if entry.get("in_stock") else "под заказ",
            "примечание": "добавлено автоматически из прайса Техносистемса"}


def refresh(rows: list[dict], sheet: list[dict], stock: list[dict],
            overrides: list[dict] | None = None) -> tuple[list[dict], dict]:
    """Новые цены и наличие строк реестра. Возвращает (строки, отчёт).

    Цена — из свежего поста (overrides), иначе из таблицы; в таблице цены нет —
    строка без цены, карточку сверка не трогает. Наличие — по складскому срезу
    (остаток > 0 — «есть»), без среза — по колонке «НАЛИЧИЕ» таблицы.
    """
    sheet, stock, posts = _keyed(list(sheet)), _keyed(list(stock)), _keyed(list(overrides or []))
    used, result, changed = set(), [], []

    def row_key(row):
        series = row.get("серия", "").lower()
        inverter = "inverter" in series and "мульти" not in series
        return model_key(row["модель"] + (" inv" if inverter else ""))

    flags: dict[str, set] = {}
    for row in rows:
        key, inverter = row_key(row)
        flags.setdefault(key, set()).add(inverter)
    twins = {key for key, seen in flags.items() if len(seen) > 1}
    for row in rows:
        fresh = dict(row)
        key, inverter = row_key(row)
        strict = key in twins
        multi = "мульти" in row.get("категория", "")
        candidates = _find(key, inverter, sheet, strict=strict)
        for item in candidates:
            used.add(id(item))
        priced = [i for i in candidates if i["prices"]]
        source = priced[0] if priced else (candidates[0] if candidates else None)
        post = next(iter(p for p in _find(key, inverter, posts, strict=strict) if p["prices"]), None)
        price_from = post or source
        prices = price_from["prices"] if price_from else []
        fresh["цена_1"] = _money(prices[0]) if prices else ""
        fresh["цена_2"] = _money(prices[1]) if len(prices) > 1 else ""
        # В остатках к коду бывает приписана буква версии: «RK-12MTA4G» = «RK-12MTA4».
        on_hand = _find(key, inverter, stock, multi_ok=multi, strict=strict, min_tail=1)
        if post is not None:
            available = post["in_stock"]
        elif stock:
            available = sum(i["qty"] for i in on_hand) > 0
        else:
            available = bool(source and source["in_stock"])
        fresh["наличие"] = "есть" if available else "под заказ"
        if (fresh["цена_1"], fresh["цена_2"], fresh["наличие"]) != (
                row.get("цена_1", ""), row.get("цена_2", ""), row.get("наличие", "")):
            changed.append(fresh)
        result.append(fresh)
    new = []
    known = {model_key(r["модель"])[0] for r in result}
    for item in sheet:
        if id(item) in used or not item["prices"] or item["_key"] in known or not _brand(item):
            continue
        known.add(item["_key"])
        row = _new_row(item)
        result.append(row)
        new.append(row)
    return result, {"changed": changed, "new": [{"model": r["модель"], "brand": r["бренд"],
                                                 "price": r["цена_1"]} for r in new]}


# ── Файлы, состояние, запуск ─────────────────────────────────────────────────


def read_csv(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with io.open(path, encoding="utf-8-sig", newline="") as handle:
        return [dict(row) for row in csv.DictReader(handle, delimiter=";")]


def write_csv(path: Path, rows: list[dict]) -> None:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=COLUMNS, delimiter=";", lineterminator="\n",
                            extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow({column: row.get(column, "") for column in COLUMNS})
    path.write_text(buffer.getvalue(), encoding="utf-8")


def load_state(path: Path = STATE_FILE) -> dict:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def fetch_sheet(gid: str) -> str:
    response = requests.get(SHEET_URL % (SHEET_KEY, gid), timeout=60)
    response.raise_for_status()
    response.encoding = "utf-8"
    return response.text


def _xlsx_rows(data: bytes) -> list[list]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return []
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    return [list(row) for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True)]


def read_posts(state: dict, posts: list[dict], now: str) -> tuple[list[dict], list[str]]:
    """Новые посты канала → свежие цены (в state["overrides"]) и строки для сообщения."""
    notes, last = [], int(state.get("tg_last_id") or 0)
    overrides = state.setdefault("overrides", [])
    for post in sorted(posts, key=lambda p: p["id"]):
        if post["id"] <= last:
            continue
        last = post["id"]
        entries = parse_post(post.get("text") or "")
        name = (post.get("file_name") or "").lower()
        if post.get("data_b64") and name.endswith((".xlsx", ".xlsm")):
            entries += parse_table_rows(_xlsx_rows(base64.b64decode(post["data_b64"])))
        priced = [e for e in entries if e["prices"]]
        when = post.get("date", now)[:16].replace("T", " ")
        if len(priced) >= 3:
            for entry in entries:
                entry["date"] = post.get("date", now)
            overrides[:] = [o for o in overrides if model_key(o["model"]) not in
                            {model_key(e["model"]) for e in entries}] + entries
            notes.append("📨 Техносистемс прислал прайс (%s): цен %d — применены" % (when, len(priced)))
        elif post.get("media") == "photo" and not entries:
            notes.append("🖼 Техносистемс прислал картинку (%s) — цены с картинки не разбираю, "
                         "сверка идёт по их таблице" % when)
        elif post.get("media") == "document":
            notes.append("📎 Техносистемс прислал файл %s (%s)%s" % (
                post.get("file_name") or "", when, " — разобран" if priced else " — цен в нём не нашёл"))
    state["tg_last_id"] = last
    return overrides, notes


def run(dry_run: bool = False, csv_path: Path | None = None) -> dict:
    csv_path = Path(csv_path or sources_techno.DEFAULT_PRICE_FILE)
    state = load_state()
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    price_text, stock_text = fetch_sheet(PRICE_GID), fetch_sheet(STOCK_GID)
    sheet = parse_sheet(price_text)
    if len([e for e in sheet if e["prices"]]) < 20:
        raise RuntimeError("таблица Техносистемса пустая или сменила вид — CSV не трогаю")
    stock_date, stock = parse_stock(stock_text)
    digest = hashlib.sha256(price_text.encode("utf-8")).hexdigest()
    if state.get("sheet_hash") != digest:
        # Таблица изменилась — она свежее всех прежних постов.
        state["sheet_hash"], state["sheet_changed_at"] = digest, now
        state["overrides"] = []
    posts = json.loads(POSTS_FILE.read_text(encoding="utf-8")) if POSTS_FILE.exists() else []
    if not state.get("tg_seeded"):
        # Канал читаем впервые: старые посты уже учтены в таблице — берём только новые.
        if posts:
            state["tg_last_id"] = max(p["id"] for p in posts)
            state["tg_seeded"] = True
        posts = []
    overrides, notes = read_posts(state, posts, now)
    before = read_csv(csv_path)
    rows, report = refresh(before, sheet, stock, overrides)
    announced = set(state.get("announced") or [])
    fresh = [row for row in report["new"] if row["model"] not in announced]
    state["announced"] = sorted(announced | {row["model"] for row in report["new"]})
    report.update(notes=notes, stock_date=stock_date, overrides=len(overrides), fresh_new=fresh)
    if not dry_run:
        write_csv(csv_path, rows)
        STATE_FILE.parent.mkdir(parents=True, exist_ok=True)
        STATE_FILE.write_text(json.dumps(state, ensure_ascii=False, indent=1), encoding="utf-8")
    return report


def summary(report: dict) -> str:
    lines = ["Техносистемс: таблица, склад от %s; изменилось строк: %d; новых позиций: %d%s" % (
        report.get("stock_date") or "—", len(report["changed"]), len(report["new"]),
        "; цены из свежего поста: %d" % report["overrides"] if report.get("overrides") else "")]
    lines += report.get("notes", [])
    for row in report["new"][:10]:
        lines.append("  ➕ %s %s — %s ₽ (карточка появится с описанием и фото)" % (
            row["brand"], row["model"], row["price"]))
    return "\n".join(lines)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Прайс Техносистемса из их таблицы и Telegram-канала")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--quiet", action="store_true", help="не слать сообщение в Telegram")
    parser.add_argument("--last-id", action="store_true", help="напечатать последний прочитанный пост")
    args = parser.parse_args(argv)
    if args.last_id:
        print(int(load_state().get("tg_last_id") or 0))
        return 0
    report = run(dry_run=args.dry_run)
    text = summary(report)
    sys.stdout.buffer.write((text + "\n").encode("utf-8"))
    # В Telegram — только новости: новый пост канала или позиции, о которых ещё не писали.
    if not args.dry_run and not args.quiet and (report["fresh_new"] or report["notes"]):
        from tools.price_sync import notify
        notify.send(text)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
